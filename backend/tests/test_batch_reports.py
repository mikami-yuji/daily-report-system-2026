import pytest
import os
import sys
import tempfile
import openpyxl
from fastapi import BackgroundTasks

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

import models
import routes_reports
import config

def test_batch_report_create_input_model():
    """BatchReportCreateInputモデルのバリデーションテスト"""
    reports_data = [
        {"日付": "26/08/18", "訪問先名": "テスト得意先A", "行動内容": "訪問（商談）"},
        {"日付": "26/08/18", "訪問先名": "テスト得意先B", "行動内容": "電話"}
    ]
    batch_input = models.BatchReportCreateInput(
        reports=[models.ReportInput(**r) for r in reports_data]
    )
    assert len(batch_input.reports) == 2
    assert batch_input.reports[0].訪問先名 == "テスト得意先A"
    assert batch_input.reports[1].訪問先名 == "テスト得意先B"

def test_add_batch_reports_empty(tmp_path):
    """空リスト時のレスポンス検証"""
    orig_db = config.SQLITE_CACHE_DB
    config.SQLITE_CACHE_DB = str(tmp_path / "test_empty.db")
    try:
        batch_input = models.BatchReportCreateInput(reports=[])
        res = routes_reports.add_batch_reports(batch_input, BackgroundTasks(), "non_existent.xlsx")
        assert res["count"] == 0
        assert res["management_numbers"] == []
    finally:
        config.SQLITE_CACHE_DB = orig_db

def test_add_batch_reports_offline_queueing(tmp_path):
    """ファイルが存在しない場合の安全退避キュー登録テスト"""
    orig_db = config.SQLITE_CACHE_DB
    config.SQLITE_CACHE_DB = str(tmp_path / "test_queue.db")
    try:
        reports_data = [
            models.ReportInput(日付="26/08/18", 訪問先名="テスト得意先A", 行動内容="訪問（商談）"),
            models.ReportInput(日付="26/08/18", 訪問先名="テスト得意先B", 行動内容="電話")
        ]
        batch_input = models.BatchReportCreateInput(reports=reports_data)
        res = routes_reports.add_batch_reports(batch_input, BackgroundTasks(), "non_existent_file_xyz.xlsx")
        assert res["offline"] is True
        assert res["count"] == 2
        assert len(res["task_ids"]) == 2
        assert len(res["management_numbers"]) == 2
        assert res["management_numbers"][0] < 0
    finally:
        config.SQLITE_CACHE_DB = orig_db

def test_apply_batch_add_reports_to_excel():
    """Excelファイルへの一括書き込みテスト"""
    wb = openpyxl.Workbook()
    ws_report = wb.active
    ws_report.title = "営業日報"
    # ヘッダー設定
    headers = [
        "管理番号", "日付", "行動内容", "エリア", "得意先CD.", "直送先CD.",
        "訪問先名", "直送先名", "重点顧客", "ランク", "得意先目標", "面談者",
        "滞在時間", "デザイン提案有無", "デザイン種別", "デザイン名", "デザイン進捗状況",
        "デザイン依頼No.", "商談内容", "提案物", "次回プラン", "競合他社情報",
        "上長コメント", "コメント返信欄", "上長", "山澄常務", "岡本常務", "中野次長",
        "既読チェック", "システム確認用デザインNo."
    ]
    ws_report.append(headers)
    # 初期データ行 (管理番号 1)
    sample_row = [1, "26/08/17", "訪問", "東京", "1001", "", "既存先", "", "重点", "A", "目標値", "担当者", "30分"] + [""] * 17
    ws_report.append(sample_row)

    ws_cust = wb.create_sheet(title="得意先_List")
    ws_cust.append(["得意先CD.", "直送先CD.", "得意先名", "直送先名", "フリガナ", "都道府県", "担当者", "重点顧客", "担当者2", "現目標"])
    ws_cust.append(["1001", "", "既存先", "", "キゾン", "東京", "営業A", "重点", "営業A", "現目標テスト値"])

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        tmp_path = tmp.name
    
    try:
        wb.save(tmp_path)
        wb.close()

        reports_to_add = [
            models.ReportInput(日付="26/08/18", 得意先CD="1001", 訪問先名="既存先", 行動内容="訪問（商談）", 面談者="山田様"),
            models.ReportInput(日付="26/08/18", 得意先CD="9999", 訪問先名="新規先", 行動内容="量販店調査", 商談内容="調査メモ")
        ]

        mgmt_nums = routes_reports._apply_batch_add_reports_to_excel(tmp_path, reports_to_add)
        assert len(mgmt_nums) == 2
        assert mgmt_nums[0] == 2
        assert mgmt_nums[1] == 3

        # 結果を再読み込みして検証
        wb_check = openpyxl.load_workbook(tmp_path)
        ws_check = wb_check["営業日報"]
        assert ws_check.max_row == 4
        # 1行目: ヘッダー, 2行目: 管理番号1, 3行目: 管理番号2, 4行目: 管理番号3
        assert ws_check.cell(row=3, column=1).value == 2
        assert ws_check.cell(row=3, column=5).value == "1001"
        assert ws_check.cell(row=3, column=7).value == "既存先"
        assert ws_check.cell(row=3, column=11).value == "現目標テスト値" # 目標がマスタから入っていること

        assert ws_check.cell(row=4, column=1).value == 3
        assert ws_check.cell(row=4, column=5).value == "9999"
        assert ws_check.cell(row=4, column=7).value == "新規先"
        wb_check.close()
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)
